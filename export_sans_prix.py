"""Exporte les vinyles de Benoit sans prix d'achat vers un fichier Excel à compléter."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
import sync

OUT = "/Users/benoittoussaint/Desktop/Phonotek vinyles sans prix.xlsx"
owner = json.load(open(sync.CONFIG_FILE))["owner"]
rows = [r for r in sync.call("GET", "/rest/v1/albums?select=id,data&owner=eq." + owner)
        if r["data"].get("support") == "Vinyle" and not r["data"].get("wish")
        and not (r["data"].get("purchase") or {}).get("price")]
rows.sort(key=lambda r: (r["data"]["artist"].lower(), r["data"]["album"].lower()))

F = Font(name="Arial", size=11)
B = Font(name="Arial", size=11, bold=True)
G = Font(name="Arial", size=9, color="999999")
Y = PatternFill("solid", start_color="FFF4B8")
H = PatternFill("solid", start_color="E8E2D6")

wb = Workbook()
ws = wb.active
ws.title = "Vinyles à compléter"
ws.append(["Identifiant (ne pas modifier)", "Artiste", "Album", "Année", "Label",
           "Prix payé (€)", "Lieu d'achat", "Date d'achat"])
for c in ws[1]:
    c.font, c.fill = B, H
    c.alignment = Alignment(wrap_text=True, vertical="center")
for r in rows:
    d = r["data"]
    ws.append([r["id"], d["artist"], d["album"], str(d.get("year") or ""), d.get("label") or "", None, None, None])
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = F
    row[0].font = G
    for c in row[5:8]:
        c.fill = Y
    row[5].number_format = '#,##0.00 "€"'
    row[7].number_format = "DD/MM/YYYY"
for col, w in zip("ABCDEFGH", [14, 30, 42, 8, 22, 14, 24, 14]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "B2"
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="10000", allow_blank=True,
                    error="Indique un montant, par exemple 24,99", errorTitle="Prix")
ws.add_data_validation(dv)
dv.add("F2:F%d" % ws.max_row)

m = wb.create_sheet("Mode d'emploi")
for text, font in [
    ("Phonotek : vinyles sans prix d'achat", B), ("", F),
    ("Remplis uniquement les 3 colonnes surlignées en jaune (F, G et H) de l'onglet « Vinyles à compléter ».", F),
    ("Laisse vide un vinyle dont tu ne connais pas le prix : il sera simplement ignoré.", F),
    ("Ne modifie pas la colonne A (identifiant) : elle permet de retrouver le disque dans Phonotek.", F), ("", F),
    ("Exemple de ligne bien remplie :", B),
    ("Prix payé : 24,99   |   Lieu d'achat : Brocante de Chatou   |   Date d'achat : 14/09/2025", F), ("", F),
    ("Une fois rempli, enregistre le fichier sans changer son nom et dis-le à Claude.", F)]:
    m.append([text])
    m.cell(m.max_row, 1).font = font
m.column_dimensions["A"].width = 110
wb.save(OUT)
print(len(rows), "vinyles exportés vers", OUT)
